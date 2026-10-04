"""
excel_report_builder.py
------------------------
ADDITIVE module. Does not modify the existing fetch/scheduler/CSV
pipeline — it only READS the snapshot CSVs that scheduler.py's
capture_snapshot_job() already writes to
forecast/<DisplayName>/<DisplayName>_<date>_<HHMM>.csv, and produces
one new .xlsx file per pilot site.

Column format matches the reference file
"24_hrs_Kothegudem_11_Aug_2026_vasila_Xweather.xlsx" exactly:
  - first column "time"
  - then, grouped by parameter (same 5 parameters/order/label templates
    as report_builder.py's REPORT_PARAMS, reused here read-only so the
    two report formats can never drift apart): ghi, temperature,
    POA_Global, POA_Diffuse, POA_Direct — each with one column per
    capture time that fired that day, e.g. "ghi_6:45", "ghi_8:15", ...
    "POA_Diffuse 6:45" (note: report_builder.py's own templates use a
    space instead of an underscore for POA_Diffuse/POA_Direct — kept
    exactly as-is to match the existing convention, not "fixed").
  - a fixed 24h/96-row grid per day: 07:00 that day -> 06:45 the next
    day, 15-min steps — same fixed window report_builder.py uses.

For a given column/row:
  - if that capture's own snapshot file has a real value at that exact
    timestamp -> used as-is, no highlight.
  - otherwise (the row falls before that capture's own data window
    even started) -> carried forward from the nearest EARLIER capture
    that has a real value at that timestamp, and the cell is filled
    YELLOW.
  - if no earlier capture has it either -> left blank. Nothing is ever
    fabricated (no 0s, no NaN substitution, no interpolation).

Existing files this module reads but never writes to: forecast/<Site>/*.csv
(read-only), and report_builder.py's REPORT_PARAMS / _capture_label
(read-only import of pure functions/data — no filesystem coupling).
Existing files this module never touches at all: everything in data/,
report_builder.py itself, xweather_service.py, app.py. scheduler.py
only gets one new call added after its existing report line.
"""

import os
import re
from datetime import datetime, timezone, timedelta

import pandas as pd
from openpyxl.styles import PatternFill

from config import SNAPSHOT_DIR, PILOT_SITES, PILOT_SITE_DISPLAY_NAMES
from report_builder import REPORT_PARAMS, _capture_label

IST = timezone(timedelta(hours=5, minutes=30))

SNAPSHOT_FILENAME_RE = re.compile(r"_(\d{4})\.csv$")

EXCEL_OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "excel_reports")

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def _find_snapshot_files(display_name, date_obj):
    """
    Returns [(hhmm_tag, file_path), ...] sorted chronologically, for
    whichever captures actually fired and wrote a file that day. Skips
    (does not fabricate) any capture that failed/didn't run.
    """
    site_dir = os.path.join(SNAPSHOT_DIR, display_name)
    if not os.path.isdir(site_dir):
        # Case-insensitive fallback (read-only lookup) — some existing
        # snapshot folders, e.g. SIRMOUR, are cased differently from the
        # display name in config.py. Never renames or creates anything.
        if os.path.isdir(SNAPSHOT_DIR):
            for entry in os.listdir(SNAPSHOT_DIR):
                if entry.lower() == display_name.lower() and os.path.isdir(os.path.join(SNAPSHOT_DIR, entry)):
                    site_dir = os.path.join(SNAPSHOT_DIR, entry)
                    break
    if not os.path.isdir(site_dir):
        return []

    date_tag = date_obj.strftime("%Y-%m-%d")
    prefix = f"{display_name}_{date_tag}_"
    matches = []
    for fname in os.listdir(site_dir):
        if fname.startswith(prefix) and fname.endswith(".csv"):
            m = SNAPSHOT_FILENAME_RE.search(fname)
            if m:
                matches.append((m.group(1), os.path.join(site_dir, fname)))
    # 4-digit zero-padded HHMM tags -> string sort == chronological sort.
    matches.sort(key=lambda pair: pair[0])
    return matches


def _load_snapshot(file_path, value_cols):
    """Loads one snapshot CSV into a DataFrame indexed by its 'time'
    column, keeping only the requested value columns that are actually
    present in this file."""
    df = pd.read_csv(file_path)
    if "time" not in df.columns:
        return None
    df["time"] = pd.to_datetime(df["time"])
    df = df.set_index("time").sort_index()
    present = [c for c in value_cols if c in df.columns]
    return df[present] if present else None


def build_excel_report(site_key, date_obj=None):
    """
    Builds excel_reports/24_hrs_<DisplayName>_<date_tag>.xlsx for one
    pilot site, matching the confirmed reference file's exact column
    layout. Returns the output path, or None if there were no snapshots
    to build from that day.
    """
    display_name = PILOT_SITE_DISPLAY_NAMES.get(site_key, site_key)
    date_obj = date_obj or datetime.now(IST).date()

    captures = _find_snapshot_files(display_name, date_obj)
    if not captures:
        print(f"[excel_report_builder] No snapshots found for {site_key} on {date_obj} — skipping.")
        return None

    value_cols = [value_col for value_col, _ in REPORT_PARAMS]
    loaded = [(hhmm_tag, _load_snapshot(fp, value_cols)) for hhmm_tag, fp in captures]

    # Fixed 24h/96-row grid, same window report_builder.py uses:
    # 07:00 that day -> 06:45 the next day, 15-min steps.
    start = pd.Timestamp(datetime.combine(date_obj, datetime.min.time()), tz=IST) + pd.Timedelta(hours=7)
    end = start + pd.Timedelta(hours=23, minutes=45)
    target_index = pd.date_range(start=start, end=end, freq="15min")

    out_columns = {}
    fallback_mask = {}  # column_name -> set of row positions that are fallback

    for value_col, label_template in REPORT_PARAMS:
        for i, (hhmm_tag, df) in enumerate(loaded):
            label = _capture_label(hhmm_tag)
            col_name = label_template.format(label=label)
            values = []
            is_fallback = set()
            for row_pos, t in enumerate(target_index):
                own_value = None
                if df is not None and value_col in df.columns and t in df.index:
                    own_value = df.at[t, value_col]
                    if pd.isna(own_value):
                        own_value = None

                if own_value is not None:
                    values.append(round(float(own_value), 2))
                    continue

                fallback_value = None
                for j in range(i - 1, -1, -1):
                    _, earlier_df = loaded[j]
                    if earlier_df is not None and value_col in earlier_df.columns and t in earlier_df.index:
                        candidate = earlier_df.at[t, value_col]
                        if not pd.isna(candidate):
                            fallback_value = candidate
                            break
                if fallback_value is not None:
                    values.append(round(float(fallback_value), 2))
                    is_fallback.add(row_pos)
                else:
                    values.append(None)  # genuinely missing — left blank, never fabricated

            out_columns[col_name] = values
            fallback_mask[col_name] = is_fallback

    result_df = pd.DataFrame(out_columns, index=target_index)
    result_df.insert(0, "time", [t.isoformat() for t in target_index])

    os.makedirs(EXCEL_OUTPUT_DIR, exist_ok=True)
    date_tag = f"{date_obj.day}_{date_obj.strftime('%B').lower()}_{date_obj.year}"
    out_path = os.path.join(EXCEL_OUTPUT_DIR, f"24_hrs_{display_name}_{date_tag}.xlsx")

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        result_df.to_excel(writer, index=False, sheet_name="Sheet1")
        ws = writer.sheets["Sheet1"]

        col_names = list(result_df.columns)
        for col_idx, col_name in enumerate(col_names):
            if col_name == "time":
                continue
            fallback_rows = fallback_mask.get(col_name, set())
            excel_col = col_idx + 1  # 1-indexed
            for row_pos in fallback_rows:
                excel_row = row_pos + 2  # +1 for header row, +1 for 1-indexing
                ws.cell(row=excel_row, column=excel_col).fill = YELLOW_FILL

    n_fallback_cells = sum(len(v) for v in fallback_mask.values())
    print(
        f"[excel_report_builder] Wrote {out_path} "
        f"({len(captures)} captures, {len(result_df)} rows, "
        f"{len(REPORT_PARAMS)} parameters, {n_fallback_cells} fallback cells highlighted)"
    )
    return out_path


def build_excel_reports_for_all_pilot_sites(date_obj=None):
    """Mirrors report_builder.build_reports_for_all_pilot_sites(), but for
    the new Excel output. Failures for one site don't stop the others."""
    for site_key in PILOT_SITES:
        try:
            build_excel_report(site_key, date_obj)
        except Exception as e:
            print(f"[excel_report_builder] FAILED for {site_key}: {type(e).__name__}: {e}")


if __name__ == "__main__":
    # Manual run: `python excel_report_builder.py` from the backend/ folder
    # builds today's combined Excel for both pilot sites.
    build_excel_reports_for_all_pilot_sites()